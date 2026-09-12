# -*- coding: utf-8 -*-
"""
pafta_report.py — JSON / Excel / CSV cikti yazicilari.
Hem pafta_ayikla.py (AutoCAD) hem pafta_dxf.py tarafindan kullanilir.
"""

from __future__ import annotations

import csv
import json
import os
from typing import Iterable, Sequence

BOM_COLS = [
    ("dosya", "Dosya"),
    ("pafta", "Pafta (Resim No)"),
    ("poz", "Poz"),
    ("resim_no", "Resim No / Ad"),
    ("adet", "Adet"),
    ("tanim", "Tanimi"),
    ("std", "Std"),
    ("malzeme", "Malzeme"),
    ("notlar", "Notlar"),
    ("birim_agirlik", "Birim Ag. [kg]"),
    ("toplam_agirlik", "Toplam Ag. [kg]"),
]

SHEET_COLS = [
    ("dosya", "Dosya"),
    ("sira", "Sira"),
    ("resim_no", "Resim No"),
    ("proje", "Proje"),
    ("parca", "Parca Adi"),
    ("is_adi", "Is / Job Name"),
    ("is_no", "Is No"),
    ("musteri", "Firma"),
    ("pafta_no", "Pafta No"),
    ("kagit", "Kagit"),
    ("olcek", "Olcek (kullanilan)"),
    ("olcek_antet", "Olcek (antette yazan)"),
    ("olcek_uyum", "Olcek uyumu"),
    ("olcek_sapma", "Olcek sapmasi [%]"),
    ("olcek_std", "Cerceve standart mi"),
    ("cerceve_w", "Cerceve G"),
    ("cerceve_h", "Cerceve Y"),
    ("kirpma_mm", "Kirpma [mm]"),
    ("malzeme_satiri", "Malzeme satiri"),
    ("toplam_agirlik", "Toplam Agirlik"),
    ("pdf", "PDF"),
    ("durum", "Durum"),
    ("uyarilar", "Uyarilar"),
]


def sheet_dict(sh, dosya: str, pdf_name: str = "", durum: str = "") -> dict:
    olcek_uyum = "-"
    if sh.scale_text_n is not None:
        olcek_uyum = "OK" if abs(sh.scale_text_n - sh.scale_n) / sh.scale_n <= 0.01 else "UYUSMUYOR"
    elif sh.scale_text:
        olcek_uyum = "OKUNAMADI"
    return {
        "dosya": dosya,
        "sira": sh.index,
        "resim_no": sh.drawing_no,
        "proje": sh.project,
        "parca": sh.part_name,
        "is_adi": sh.job_name,
        "is_no": sh.project_no,
        "musteri": sh.customer,
        "pafta_no": sh.sheet_no,
        "kagit": sh.paper,
        "olcek": sh.scale_label,
        "olcek_antet": sh.scale_text,
        "olcek_uyum": olcek_uyum,
        "olcek_sapma": getattr(sh, "scale_deviation", 0.0),
        "olcek_std": "EVET" if getattr(sh, "scale_is_standard", True) else "HAYIR",
        "cerceve_w": round(sh.frame.w, 2),
        "cerceve_h": round(sh.frame.h, 2),
        "kirpma_mm": sh.crop_mm,
        "malzeme_satiri": len(sh.bom),
        "toplam_agirlik": sh.bom_total_weight,
        "pdf": pdf_name,
        "durum": durum,
        "uyarilar": " | ".join(sh.warnings),
        "pencere": {"ll": list(sh.window_ll), "ur": list(sh.window_ur)},
        "malzeme_listesi": sh.bom,
    }


def write_json(path: str, sheets: Sequence[dict]) -> None:
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump({"pafta_sayisi": len(sheets), "paftalar": list(sheets)},
                  f, ensure_ascii=False, indent=2)


def write_csv(path: str, sheets: Sequence[dict]) -> None:
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([h for _, h in SHEET_COLS])
        for s in sheets:
            w.writerow([s.get(k, "") for k, _ in SHEET_COLS])


def write_bom_csv(path: str, rows: Sequence[dict]) -> None:
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([h for _, h in BOM_COLS])
        for r in rows:
            w.writerow([r.get(k, "") for k, _ in BOM_COLS])


def write_xlsx(path: str, sheets: Sequence[dict], rows: Sequence[dict]) -> bool:
    """Tek Excel dosyasi: 'Malzeme Listesi' + 'Paftalar' sayfalari.
    openpyxl yoksa CSV'ye duser ve False doner."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        base = os.path.splitext(path)[0]
        write_bom_csv(base + "_malzeme.csv", rows)
        write_csv(base + "_paftalar.csv", sheets)
        print("  ! openpyxl kurulu degil, Excel yerine CSV yazildi "
              "(pip install openpyxl).")
        return False

    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(color="FFFFFF", bold=True)

    def dump(ws, cols, data, numeric_keys=()):
        ws.append([h for _, h in cols])
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for d in data:
            vals = []
            for k, _ in cols:
                v = d.get(k, "")
                if k in numeric_keys and isinstance(v, str):
                    t = v.replace(".", "").replace(",", ".") if v.count(",") == 1 and v.count(".") > 1 else v.replace(",", ".")
                    try:
                        v = float(t)
                    except (TypeError, ValueError):
                        pass
                vals.append(v)
            ws.append(vals)
        ws.freeze_panes = "A2"
        widths = []
        for i, (k, h) in enumerate(cols, 1):
            m = len(str(h))
            for d in data:
                m = max(m, min(48, len(str(d.get(k, "")))))
            widths.append(m + 2)
            ws.column_dimensions[get_column_letter(i)].width = widths[-1]
        if len(data) > 0:
            ref = f"A1:{get_column_letter(len(cols))}{len(data)+1}"
            t = Table(displayName=f"T_{ws.title.replace(' ', '')}", ref=ref)
            t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
            try:
                ws.add_table(t)
            except Exception:
                pass

    ws1 = wb.active
    ws1.title = "Malzeme Listesi"
    dump(ws1, BOM_COLS, rows, numeric_keys=("adet", "birim_agirlik", "toplam_agirlik"))

    ws2 = wb.create_sheet("Paftalar")
    dump(ws2, SHEET_COLS, sheets)

    wb.save(path)
    return True
