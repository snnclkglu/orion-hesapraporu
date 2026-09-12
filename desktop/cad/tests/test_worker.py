import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from types import SimpleNamespace
from unittest.mock import Mock
import queue
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import worker


class WorkerSafetyTests(unittest.TestCase):
    def test_paired_worker_monitors_and_wakes_on_status_change(self):
        events = queue.Queue()
        w = worker.Worker({'origin':'https://example.com'}, events)
        self.assertTrue(w.active.is_set())
        w.status('ready','Hazır')
        self.assertTrue(w.heartbeat_wake.is_set())
        self.assertEqual(events.get_nowait(), 'Hazır')
        w.heartbeat_wake.clear()
        w.status('ready','Hazır')
        self.assertFalse(w.heartbeat_wake.is_set())
        self.assertTrue(events.empty())

    def test_open_drawing_wait_recovers_without_manual_restart(self):
        w = worker.Worker({'origin':'https://example.com'}, queue.Queue())
        w.stop = Mock()
        w.stop.is_set.side_effect = [False, False, False, True]
        w.api.call = Mock(side_effect=[{'accepted':True}, {'job':None}])
        with patch.object(worker.threading, 'Thread'), patch.object(worker, 'autocad_probe', side_effect=[('attention','','Çizim açık'),('ready','25','Hazır')]):
            w.run()
        self.assertTrue(w.active.is_set())
        self.assertEqual(w.api.call.call_args_list[-1].args, ('claim',))
        self.assertEqual(w.state, 'ready')

    def test_plot_proxy_retries_rejected_properties_and_methods(self):
        error = RuntimeError('busy'); error.hresult = -2147418111
        class Layout:
            def __init__(self): self.writes = 0; self.reads = 0
            @property
            def ConfigName(self):
                self.reads += 1
                if self.reads == 1: raise error
                return 'PDF.pc3'
            @ConfigName.setter
            def ConfigName(self, value):
                self.writes += 1
                if self.writes == 1: raise error
        target = Layout(); target.PlotToFile = Mock(side_effect=[error,True])
        proxy = worker.PlotComProxy(target)
        with patch.object(worker.time, 'sleep'):
            self.assertEqual(proxy.ConfigName, 'PDF.pc3')
            proxy.ConfigName = 'PDF.pc3'
            self.assertTrue(proxy.PlotToFile('test.pdf'))
        self.assertEqual(target.writes,2)
        self.assertEqual(target.PlotToFile.call_count,2)

    def test_plot_proxy_does_not_retry_unknown_or_completed_calls(self):
        target = SimpleNamespace(PlotToFile=Mock(return_value=True))
        worker.PlotComProxy(target).PlotToFile('test.pdf')
        self.assertEqual(target.PlotToFile.call_count,1)
        target.PlotToFile = Mock(side_effect=RuntimeError('invalid plotter'))
        with self.assertRaises(RuntimeError): worker.PlotComProxy(target).PlotToFile('test.pdf')
        self.assertEqual(target.PlotToFile.call_count,1)

    def test_plot_wait_requires_idle_command_and_has_deadline(self):
        acad = self.acad([],idle=True)
        doc = SimpleNamespace(GetVariable=Mock(side_effect=[1,0]))
        with patch.object(worker.time,'sleep'):
            worker.wait_plot_ready(acad,doc)
        self.assertEqual(doc.GetVariable.call_count,2)
        with patch.object(worker.time,'monotonic',side_effect=[0,1,31]), patch.object(worker.time,'sleep'):
            with self.assertRaises(RuntimeError): worker.wait_plot_ready(self.acad([],idle=False),doc)

    def test_repair_replaces_credentials_only_after_valid_pair(self):
        existing = Mock()
        order = []
        existing.stop_for_pairing.side_effect = lambda: order.append('stopped')
        with patch.object(worker.Api, 'call', side_effect=lambda *a, **k: (order.append('paired') or {'deviceId':'new-device','storageOrigin':'https://data.supabase.co'})), patch.object(worker, 'save_credentials', side_effect=lambda value: order.append('saved')):
            value = worker.pair_connection(existing, 'https://example.com', 'a'*64)
        self.assertEqual(order, ['stopped','paired','saved'])
        self.assertEqual(value['deviceId'], 'new-device')

    def test_repair_failure_never_overwrites_credentials(self):
        existing = Mock()
        with patch.object(worker, 'save_credentials') as save, patch.object(worker.Api, 'call', side_effect=worker.ApiError('expired',400)):
            with self.assertRaises(ValueError): worker.pair_connection(existing,'https://example.com','bad')
            existing.stop_for_pairing.assert_not_called()
            with self.assertRaises(worker.ApiError): worker.pair_connection(existing,'https://example.com','a'*64)
            save.assert_not_called()

    def test_repair_blocks_active_work_and_waits_for_old_threads(self):
        w = worker.Worker({'origin':'https://example.com'}, queue.Queue())
        w.active.set()
        with self.assertRaises(RuntimeError): w.stop_for_pairing()
        self.assertFalse(w.stop.is_set())
        w.active.clear(); w.job = {'id':'test'}
        with self.assertRaises(RuntimeError): w.stop_for_pairing()
        w.job = None
        w.thread = Mock(); w.thread.is_alive.return_value = True
        with self.assertRaises(RuntimeError): w.stop_for_pairing()
        w.thread.is_alive.return_value = False
        w.stop_for_pairing()
        self.assertTrue(w.stop.is_set())

    def collection(self, values):
        return SimpleNamespace(Count=len(values), Item=lambda i: values[i])

    def blank(self, titled=0, modified=0, path='', model=0, entities=('AcDbViewport',)):
        return SimpleNamespace(Path=path, GetVariable=lambda key: {'DWGTITLED':titled,'DBMOD':modified}[key],
                               ModelSpace=SimpleNamespace(Count=model),
                               Layouts=self.collection([SimpleNamespace(Block=self.collection([SimpleNamespace(ObjectName=e) for e in entities]))]))

    def acad(self, docs, idle=True):
        return SimpleNamespace(Documents=SimpleNamespace(Count=len(docs),Item=lambda i:docs[i]),
                               GetAcadState=lambda:SimpleNamespace(IsQuiescent=idle))

    def test_pristine_blank_can_remain_open(self):
        worker.assert_autocad_ready(self.acad([]))
        worker.assert_autocad_ready(self.acad([self.blank()]))
        worker.assert_autocad_ready(self.acad([self.blank(path='C:/WINDOWS/system32')]))

    def test_real_modified_and_nonempty_drawings_are_blocked(self):
        for options in ({'titled':1},{'modified':1},{'modified':4},
                        {'model':1},{'entities':('AcDbViewport','AcDbLine')}):
            with self.subTest(options=options):
                with self.assertRaises(RuntimeError):
                    worker.assert_autocad_ready(self.acad([self.blank(),self.blank(**options)]))

    def test_unknown_document_and_busy_autocad_are_blocked(self):
        with self.assertRaises(RuntimeError):
            worker.assert_autocad_ready(self.acad([SimpleNamespace()]))
        with self.assertRaises(RuntimeError):
            worker.assert_autocad_ready(self.acad([self.blank()],idle=False))

    def test_temporary_rpc_rejection_is_retried_but_bounded(self):
        from unittest.mock import Mock
        error = RuntimeError('busy'); error.hresult = -2147418111
        read = Mock(side_effect=[error, 42])
        with patch.object(worker.time, 'sleep'):
            self.assertEqual(worker.com_read(read), 42)
            always_busy = Mock(side_effect=error)
            with self.assertRaises(RuntimeError): worker.com_read(always_busy)
        self.assertEqual(always_busy.call_count, 21)

    def test_windows_names(self):
        for name in ('../a.pdf', 'a/b.pdf', 'NUL.dwg', 'a:stream', 'a.pdf.', ' a.dwg', 'a\0.pdf'):
            self.assertFalse(worker.safe_name(name), name)
        self.assertTrue(worker.safe_name('KÖPRÜ YÜRÜTME.pdf'))

    def test_origin_rejects_credentials_and_external_http(self):
        for url in ('http://example.com', 'https://user:pass@example.com', 'https://example.com/path', 'file:///c:/a', 'https://example.com?x=y'):
            with self.assertRaises(ValueError):
                worker.validate_origin(url)
        self.assertEqual(worker.validate_origin('https://example.com/'), 'https://example.com')
        self.assertEqual(worker.validate_origin('http://localhost:3000'), 'http://localhost:3000')

    def test_storage_origin_is_bound_and_bearer_not_shared(self):
        api = worker.Api({'origin': 'https://example.com', 'storageOrigin': 'https://data.supabase.co', 'token': 'secret'})
        self.assertTrue(api.storage_url('https://data.supabase.co/a').endswith('/a'))
        for url in ('https://other.example/a', 'http://data.supabase.co/a', 'https://user@data.supabase.co/a'):
            with self.assertRaises(ValueError):
                api.storage_url(url)

    def make_output(self, root, pdf=True, errors=0, repeated=False):
        folder = root / 'output'; folder.mkdir()
        if pdf:
            (folder / 'A.pdf').write_bytes(b'%PDF-1.7\nfixture')
        sheets = [{'pdf': 'A.pdf'}] * (2 if repeated else 1)
        value = {'arac_surum': 'test', 'ozet': {'pafta': len(sheets), 'pdf_basarili': len(sheets), 'hata': errors}, 'paftalar': sheets, 'cikti_klasoru': str(folder)}
        report = root / 'raw.json'; report.write_text(json.dumps(value), encoding='utf-8')
        return folder, report

    def test_outputs_require_all_distinct_real_pdfs(self):
        for options in ({'pdf': False}, {'errors': 1}, {'repeated': True}):
            with tempfile.TemporaryDirectory() as tmp:
                folder, report = self.make_output(Path(tmp), **options)
                with self.assertRaises(ValueError):
                    worker.output_files(folder, report)

    def test_result_strips_local_path_and_is_repeatable(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder, report = self.make_output(Path(tmp))
            first = worker.output_files(folder, report)
            second = worker.output_files(folder, report)
            self.assertEqual([p.name for p, _ in first], [p.name for p, _ in second])
            self.assertNotIn('cikti_klasoru', json.loads((folder/'result.json').read_text(encoding='utf-8')))

    def test_hash_covers_content_and_atomic_json(self):
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp)/'state.json'
            worker.save_json(target, {'name': 'KÖPRÜ'})
            a = worker.digest(target)
            worker.save_json(target, {'name': 'BAŞKİRİŞ'})
            self.assertNotEqual(a, worker.digest(target))
            self.assertFalse(target.with_suffix('.json.tmp').exists())

    def test_export_workbook_preserves_values_and_never_executes_formulas(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            target = worker.write_export_workbook(Path(tmp), {"malzeme": [{"resim_no": "0026-01-0110", "tanim": "=1+1", "adet": "4", "birim_agirlik": "1,25", "toplam_agirlik": "5", "pafta": "0026-01-0100"}]})
            book = load_workbook(target)
            self.assertEqual(book.sheetnames, ["BOM"])
            sheet = book.active
            self.assertEqual(sheet['A1'].value, 'Part Number')
            self.assertEqual(sheet['B2'].value, '=1+1')
            self.assertEqual(sheet['B2'].data_type, 's')
            self.assertEqual(sheet['D2'].value, '4')
            self.assertEqual(sheet['E2'].value, '1,25')
            book.close()

    def test_no_global_autocad_termination(self):
        source = Path(worker.__file__).read_text(encoding='utf-8')
        self.assertNotIn('["taskkill"', source)
        self.assertNotIn('.Quit(', source)
        self.assertIn('process.terminate()', source)
        self.assertIn('assert_autocad_ready(application)', source)


if __name__ == '__main__':
    unittest.main()
