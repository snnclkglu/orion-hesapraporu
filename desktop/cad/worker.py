"""ORION Çizim İşleme yardımcısı. Ağ işleri AutoCAD'in sahibi olan yerel oturumda yürür."""
from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import queue
import re
import secrets
import subprocess
import sys
import threading
import time
from urllib.parse import urlparse
import uuid

VERSION = "1.0.2"
PROTOCOL = 1
MAX_BYTES = 100 * 1024 * 1024
MAX_RESULT = 2 * 1024 * 1024
TIMEOUT = 20 * 60


def app_dir() -> Path:
    root = Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "OrionCad"
    root.mkdir(parents=True, exist_ok=True)
    return root


def safe_name(value: str) -> bool:
    return (isinstance(value, str) and 0 < len(value) <= 180 and value == value.strip()
            and not re.search(r'[<>:"/\\|?*\x00-\x1f]', value) and not value.endswith(".")
            and value not in (".", "..") and not re.match(r"^(con|prn|aux|nul|com[1-9]|lpt[1-9])(?:\.|$)", value, re.I))


def validate_origin(value: str) -> str:
    parsed = urlparse(value.strip())
    local = parsed.hostname in ("localhost", "127.0.0.1", "::1")
    if (parsed.scheme != "https" and not (local and parsed.scheme == "http")) or not parsed.hostname:
        raise ValueError("HTTPS uygulama adresi gerekli.")
    if parsed.username or parsed.password or parsed.query or parsed.fragment or parsed.path not in ("", "/"):
        raise ValueError("Yalnız uygulamanın ana adresini girin.")
    return f"{parsed.scheme}://{parsed.netloc}"


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def save_json(path: Path, value: dict) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def save_credentials(value: dict) -> None:
    import win32crypt
    blob = win32crypt.CryptProtectData(json.dumps(value).encode(), "ORION CAD", None, None, None, 0)
    path = app_dir() / "connection.dat"
    tmp = path.with_suffix(".tmp")
    tmp.write_bytes(blob)
    tmp.replace(path)


def load_credentials() -> dict:
    import win32crypt
    path = app_dir() / "connection.dat"
    if not path.exists():
        return {}
    return json.loads(win32crypt.CryptUnprotectData(path.read_bytes(), None, None, None, 0)[1])


class ApiError(RuntimeError):
    def __init__(self, message: str, status: int):
        super().__init__(message)
        self.status = status


class Api:
    def __init__(self, config: dict):
        self.config = config
        self.origin = validate_origin(config["origin"])

    def call(self, action: str, data: dict | None = None, paired: bool = True) -> dict:
        import requests
        headers = {"Content-Type": "application/json"}
        if paired:
            headers["Authorization"] = "Bearer " + self.config["token"]
        response = requests.post(self.origin + "/api/cad/worker", headers=headers,
                                 json={"action": action, "data": data or {}}, timeout=(15, 90), allow_redirects=False)
        if not response.ok:
            try:
                message = response.json().get("error", "Hizmet yanıt vermedi.")
            except Exception:
                message = "Uygulama adresi veya bağlantı kontrol edilmeli."
            raise ApiError(str(message), response.status_code)
        if response.status_code != 200:
            raise ApiError("Uygulama yönlendirmesi kabul edilmedi.", response.status_code)
        return response.json()

    def storage_url(self, value: str) -> str:
        parsed = urlparse(value)
        expected = urlparse(self.config["storageOrigin"])
        if parsed.scheme != "https" or parsed.netloc != expected.netloc or parsed.username or parsed.password:
            raise ValueError("Depo adresi beklenen sunucuya ait değil.")
        return value

    def download(self, job: dict, target: Path) -> None:
        import requests
        expected = int(job["source_size"])
        if not 0 < expected <= MAX_BYTES:
            raise ValueError("DWG boyutu desteklenmiyor.")
        if target.exists() and target.stat().st_size == expected and digest(target) == job["source_sha256"]:
            return
        tmp = target.with_suffix(".part")
        size = 0
        with requests.get(self.storage_url(job["downloadUrl"]), stream=True, timeout=(15, 60), allow_redirects=False) as response:
            response.raise_for_status()
            if response.status_code != 200:
                raise ValueError("DWG indirilemedi.")
            with tmp.open("wb") as output:
                for block in response.iter_content(1024 * 1024):
                    size += len(block)
                    if size > expected:
                        raise ValueError("DWG beklenen boyutu aşıyor.")
                    output.write(block)
        if size != expected or digest(tmp) != job["source_sha256"]:
            raise ValueError("DWG bütünlük kontrolünden geçemedi.")
        tmp.replace(target)

    def upload(self, job: dict, path: Path, kind: str) -> None:
        import requests
        details = {"jobId": job["id"], "attemptId": job["attempt_id"], "name": path.name,
                   "size": path.stat().st_size, "sha256": digest(path), "kind": kind}
        target = self.call("artifact", details)["upload"]
        if target.get("exists"):
            return
        content_type = "application/pdf" if kind == "pdf" else "application/json" if path.suffix == ".json" else "application/octet-stream"
        with path.open("rb") as stream:
            response = requests.put(self.storage_url(target["signedUrl"]), data=stream,
                                    headers={"Content-Type": content_type}, timeout=(15, 180), allow_redirects=False)
        if response.status_code not in (200, 201):
            # Yanıt kaybolmuş olabilir; bir sonraki tur depo boyutunu yeniden kontrol eder.
            raise RuntimeError("Çıktı yüklenemedi; dosya yerelde korundu.")


def com_read(read):
    # AutoCAD açılışında kısa süreli RPC reddi görülebilir; yalnız okumayı tekrarlar.
    for attempt in range(21):
        try:
            return read()
        except Exception as error:
            if getattr(error, "hresult", None) not in (-2147418111, -2147417846) or attempt == 20:
                raise
            time.sleep(0.25)


def is_pristine_startup_document(document) -> bool:
    """Yalnız değiştirilmemiş, adsız, içeriksiz başlangıç çizimi; salt okunur."""
    try:
        # Path adsız Drawing1 için de çalışma klasörü olabilir; dosya varlığı ölçütü değildir.
        if int(com_read(lambda: document.GetVariable("DWGTITLED"))) != 0:
            return False
        if int(com_read(lambda: document.GetVariable("DBMOD"))) != 0 or int(com_read(lambda: document.ModelSpace.Count)) != 0:
            return False
        layouts = com_read(lambda: document.Layouts)
        for index in range(int(com_read(lambda: layouts.Count))):
            block = com_read(lambda: layouts.Item(index).Block)
            for entity_index in range(int(com_read(lambda: block.Count))):
                if str(com_read(lambda: block.Item(entity_index).ObjectName)) != "AcDbViewport":
                    return False
        return True
    except Exception:
        return False


def assert_autocad_ready(acad) -> None:
    if not com_read(lambda: acad.GetAcadState().IsQuiescent):
        raise RuntimeError("AutoCAD bir komut veya pencere bekliyor. Başlangıç/giriş pencerelerini tamamlayın.")
    documents = com_read(lambda: acad.Documents)
    for index in range(int(com_read(lambda: documents.Count))):
        document = com_read(lambda: documents.Item(index))
        if not is_pristine_startup_document(document):
            raise RuntimeError("Açık veya değiştirilmiş AutoCAD çizimlerini kaydedip kapatın; sonra kontrolü yeniden başlatın. Boş başlangıç sekmesi kalabilir.")


def autocad_probe(start: bool = False) -> tuple[str, str, str]:
    if os.name != "nt":
        return "autocad_missing", "", "Windows gerekli."
    import winreg
    try:
        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "AutoCAD.Application"):
            pass
    except OSError:
        return "autocad_missing", "", "Tam AutoCAD bulunamadı. AutoCAD LT desteklenmiyor."
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    try:
        try:
            acad = win32com.client.GetActiveObject("AutoCAD.Application")
        except Exception:
            if not start:
                return "attention", "", "AutoCAD'i açın veya Kontrol et ve başlat düğmesini kullanın."
            acad = win32com.client.Dispatch("AutoCAD.Application")
            acad.Visible = True
        version = str(acad.Version)
        try:
            assert_autocad_ready(acad)
        except RuntimeError as error:
            return "attention", version, str(error)
        return "ready", version, "Hazır. İşlem sırasında AutoCAD'de başka çizim açmayın."
    except Exception:
        return "attention", "", "AutoCAD'e erişilemiyor. Başlangıç/giriş pencerelerini tamamlayın."
    finally:
        pythoncom.CoUninitialize()


def write_export_workbook(directory: Path, value: dict) -> Path:
    # Kaynak motorun raporu korunur. Aktarım için yalnız BOM içeren ayrı kitap.
    # Adet/ağırlık yeniden hesaplanmaz; kaynak hücre değerleri aynen taşınır.
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "BOM"
    columns = [("resim_no", "Part Number"), ("tanim", "Description"),
               ("malzeme", "Material"), ("adet", "Item QTY"),
               ("birim_agirlik", "Mass"), ("pafta", "Kaynak Pafta"),
               ("poz", "Poz"), ("std", "Standart"),
               ("toplam_agirlik", "Toplam Ağırlık (kg)"), ("notlar", "Notlar")]
    ws.append([header for _, header in columns])
    for row in value.get("malzeme", []):
        ws.append([row.get(key, "") for key, _ in columns])
        # DWG metni Excel formülü olarak çalıştırılmaz.
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str): cell.data_type = "s"
    ws.freeze_panes = "A2"
    path = directory / "CAD_MALZEME.xlsx"
    wb.save(path)
    return path


def output_files(directory: Path, result_path: Path) -> list[tuple[Path, str]]:
    value = json.loads(result_path.read_text(encoding="utf-8"))
    summary = value.get("ozet", {})
    sheets = value.get("paftalar", [])
    if summary.get("hata") != 0 or not sheets or summary.get("pafta") != len(sheets) or summary.get("pdf_basarili") != len(sheets):
        raise ValueError("Tüm paftalar PDF'e dönüşmedi. Sonuçlar inceleme için yerelde korundu.")
    write_export_workbook(directory, value)
    files: dict[str, Path] = {}
    root = directory.resolve()
    for path in directory.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in (".pdf", ".xlsx", ".csv", ".json", ".txt"):
            continue
        if path.is_symlink() or not path.resolve().is_relative_to(root) or not safe_name(path.name):
            raise ValueError("Çıktı yolu desteklenmiyor.")
        if path.name in files:
            raise ValueError("Aynı isimli iki çıktı var; otomatik üzerine yazılmadı.")
        if not 0 < path.stat().st_size <= MAX_BYTES:
            raise ValueError("Çıktı boyutu desteklenmiyor.")
        files[path.name] = path
    used = set()
    for sheet in sheets:
        name = sheet.get("pdf", "")
        if not safe_name(name) or name not in files or name in used or files[name].suffix.lower() != ".pdf":
            raise ValueError("Pafta PDF'i eksik veya yinelenmiş.")
        with files[name].open("rb") as pdf:
            if pdf.read(5) != b"%PDF-":
                raise ValueError("Geçersiz PDF çıktısı.")
        used.add(name)
    # Yerel yol web sonucuna taşınmaz; bulgu verisi korunur.
    value.pop("cikti_klasoru", None)
    canonical = directory / "result.json"
    canonical.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
    if canonical.stat().st_size > MAX_RESULT:
        raise ValueError("Sonuç raporu 2 MB sınırını aşıyor.")
    outputs = [(p, "pdf" if p.suffix.lower() == ".pdf" else "diagnostic" if p.name.startswith("tani_") else "report") for p in files.values() if p.name != "result.json"]
    return outputs + [(canonical, "result")]


def run_engine(job: dict, root: Path, stop: threading.Event) -> None:
    args = ["--engine", str(root / "source.dwg"), str(root / "output"), str(root / "raw-result.json"),
            job["options"]["paper"], job["options"]["duplicates"]]
    command = [sys.executable] + ([] if getattr(sys, "frozen", False) else [str(Path(__file__).resolve())]) + args
    with (root / "engine.log").open("w", encoding="utf-8") as log:
        process = subprocess.Popen(command, stdout=log, stderr=log, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        deadline = time.monotonic() + TIMEOUT
        while process.poll() is None:
            if stop.wait(1) or time.monotonic() > deadline:
                # Yalnız kendi Python iş sürecimiz. AutoCAD'e taskkill / Quit gönderilmez.
                process.terminate()
                process.wait(timeout=15)
                raise RuntimeError("İş durdu. AutoCAD zorla kapatılmadı; açık işlem çizimini ve başlangıç ayarlarını kontrol edin.")
        if process.returncode != 0:
            raise RuntimeError("AutoCAD işlemi tamamlanamadı. Yerel iş klasöründeki engine.log dosyasını inceleyin.")


def engine_main(args: list[str]) -> int:
    source, output, result, paper, duplicates = args
    if paper not in ("A3", "AUTO") or duplicates not in ("hepsi", "alt", "ust", "dur"):
        raise ValueError("Baskı ayarları desteklenmiyor.")
    sys.path.insert(0, str(Path(__file__).parent / "engine"))
    import pythoncom
    import win32com.client
    import pafta_ayikla as engine
    pythoncom.CoInitialize()
    acad = win32com.client.GetActiveObject("AutoCAD.Application")
    assert_autocad_ready(acad)
    # Eski aracın genel önbellek silen bağlantısı kullanılmaz.
    try:
        acad = win32com.client.gencache.EnsureDispatch(acad)
    except Exception:
        pass
    engine.connect_autocad = lambda visible=True: acad
    original_open = engine.open_document
    expected = os.path.normcase(os.path.abspath(source))

    class GuardedDocument:
        def __init__(self, document):
            object.__setattr__(self, "_document", document)
            object.__setattr__(self, "_variables", {v: document.GetVariable(v) for v in ("FILEDIA", "CMDDIA", "BACKGROUNDPLOT")})

        def __getattr__(self, key):
            return getattr(self._document, key)

        def __setattr__(self, key, value):
            setattr(self._document, key, value)

        def Close(self, save=False):
            if os.path.normcase(os.path.abspath(self._document.FullName)) != expected:
                raise RuntimeError("Kullanıcı çizimi kapatılmadı.")
            for key, value in self._variables.items():
                self._document.SetVariable(key, value)
            self._document.Close(False)

    def safe_open(application, path, read_only=True):
        if os.path.normcase(os.path.abspath(path)) != expected:
            raise RuntimeError("Yalnız işin geçici DWG kopyası açılabilir.")
        assert_autocad_ready(application)
        document, opened = original_open(application, path, True)
        if not opened or os.path.normcase(os.path.abspath(document.FullName)) != expected:
            raise RuntimeError("AutoCAD beklenmeyen çizime yöneldi; çizim değiştirilmedi.")
        return GuardedDocument(document), True

    engine.open_document = safe_open
    sys.argv = ["pafta_ayikla", source, "--out", output, "--sonuc", result, "--paper", paper, "--kopya", duplicates]
    try:
        return engine.main()
    finally:
        pythoncom.CoUninitialize()


class Worker:
    def __init__(self, config: dict, events: queue.Queue):
        self.api = Api(config)
        self.events = events
        self.stop = threading.Event()
        self.active = threading.Event()
        self.job_stop = threading.Event()
        self.state = "attention"
        self.version = ""
        self.message = "Kontrol et ve başlat düğmesini kullanın."
        self.job: dict | None = None
        self.progress = ""
        self.lock = threading.Lock()

    def status(self, state: str, message: str) -> None:
        with self.lock:
            self.state, self.message = state, message
        self.events.put(message)

    def heartbeat(self) -> None:
        while not self.stop.is_set():
            with self.lock:
                body = {"state": self.state, "message": self.message[:500], "autocadVersion": self.version,
                        "helperVersion": VERSION, "protocol": PROTOCOL}
                if self.job:
                    body.update(jobId=self.job["id"], attemptId=self.job["attempt_id"], progress=self.progress)
            try:
                answer = self.api.call("heartbeat", body)
                if not answer.get("accepted", False):
                    self.job_stop.set()
                    self.active.clear()
                    self.status("attention", "İş iptal edilmiş veya bağlantı süresi dolmuş. Sonuçlar yerelde korundu.")
            except ApiError as error:
                if error.status in (401, 403):
                    self.job_stop.set(); self.active.clear()
                self.events.put(str(error))
            except Exception:
                self.events.put("İnternet bağlantısı kesildi; yerel dosyalar korunuyor.")
            self.stop.wait(25)

    def process(self, job: dict) -> None:
        job_id, attempt_id = str(uuid.UUID(job["id"])), str(uuid.UUID(job["attempt_id"]))
        root = app_dir() / "jobs" / job_id / attempt_id
        root.mkdir(parents=True, exist_ok=True)
        (root / "output").mkdir(exist_ok=True)
        # İmzalı URL ve cihaz anahtarı çalışma günlüğüne yazılmaz.
        save_json(root / "job.json", {k: v for k, v in job.items() if k != "downloadUrl"})
        self.job_stop.clear()
        with self.lock:
            self.job = job; self.progress = "DWG indiriliyor"
        self.status("busy", "AutoCAD çalışıyor. Başka çizim açmayın.")
        try:
            self.api.download(job, root / "source.dwg")
            with self.lock:
                self.progress = "Paftalar AutoCAD ile hazırlanıyor"
            run_engine(job, root, self.job_stop)
            outputs = output_files(root / "output", root / "raw-result.json")
            save_json(root / "pending.json", {"files": [{"path": str(p.relative_to(root)), "kind": kind} for p, kind in outputs]})
            for retry in range(20):
                if self.job_stop.is_set() or self.stop.is_set():
                    raise RuntimeError("İş durdu. Çıktılar yerel iş klasöründe korundu.")
                try:
                    for index, (path, kind) in enumerate(outputs):
                        if self.job_stop.is_set():
                            raise RuntimeError("İş sahipliği sona erdi.")
                        with self.lock:
                            self.progress = f"Çıktılar yükleniyor ({index+1}/{len(outputs)})"
                        self.api.upload(job, path, kind)
                    self.api.call("complete", {"jobId": job_id, "attemptId": attempt_id})
                    save_json(root / "completed.json", {"completed": True})
                    self.events.put("Tamamlandı. Sonuçlar web uygulamasında incelemeye hazır.")
                    break
                except ApiError as error:
                    if error.status in (401, 403, 409) or retry == 19:
                        raise
                    self.events.put("Yükleme yeniden denenecek; dosyalar korundu.")
                    self.stop.wait(min(5 * (retry + 1), 30))
                except Exception:
                    if retry == 19:
                        raise
                    self.events.put("Bağlantı bekleniyor; üretilen dosyalar yerelde korundu.")
                    self.stop.wait(min(5 * (retry + 1), 30))
        except Exception as error:
            self.active.clear()
            self.status("attention", str(error))
            try:
                self.api.call("fail", {"jobId": job_id, "attemptId": attempt_id, "message": str(error)[:1500]})
            except Exception:
                pass
        finally:
            with self.lock:
                self.job = None
            if self.active.is_set():
                self.status("ready", "Sonraki işlem bekleniyor.")
            elif self.state == "busy":
                self.status("attention", "Yeni iş alımı duraklatıldı.")

    def run(self) -> None:
        threading.Thread(target=self.heartbeat, daemon=True).start()
        while not self.stop.is_set():
            if not self.active.wait(1):
                continue
            try:
                state, version, message = autocad_probe()
                self.version = version
                self.status(state, message)
                if state != "ready":
                    self.active.clear()
                    continue
                # Claim öncesinde hazır durumu sunucuda görünür olmalı.
                self.api.call("heartbeat", {"state": state, "autocadVersion": version, "helperVersion": VERSION, "protocol": PROTOCOL, "message": message})
                job = self.api.call("claim").get("job")
                if job:
                    self.process(job)
                else:
                    self.stop.wait(10)
            except Exception as error:
                self.status("attention", str(error))
                self.active.clear()


def main() -> int:
    if "--engine" in sys.argv:
        # Pencereli EXE'de stdout yoktur; hata iletişim kutusu işi kilitlemesin.
        import traceback
        args = sys.argv[sys.argv.index("--engine") + 1:]
        log_path = Path(args[0]).parent / "engine.log"
        with log_path.open("a", encoding="utf-8") as log:
            sys.stdout = sys.stderr = log
            try:
                return engine_main(args)
            except Exception:
                traceback.print_exc()
                return 1
    if "--self-test" in sys.argv:
        import pythoncom, win32crypt, win32timezone, openpyxl, pypdf
        import tkinter
        assert safe_name("KÖPRÜ.dwg") and not safe_name("../test.dwg")
        print("ORION CAD " + VERSION + " · protokol " + str(PROTOCOL))
        return 0
    if os.name != "nt":
        raise SystemExit("Windows gerekli.")
    import tkinter as tk
    from tkinter import ttk, messagebox
    import win32event
    import win32api
    mutex = win32event.CreateMutex(None, False, "Local\\OrionCadWorker")
    if win32api.GetLastError() == 183:
        messagebox.showinfo("ORION", "Yardımcı zaten açık.")
        return 0
    events: queue.Queue = queue.Queue()
    root = tk.Tk()
    root.title("ORION · Çizim İşleme Yardımcısı")
    root.geometry("680x530")
    root.minsize(550, 480)
    panel = ttk.Frame(root, padding=24); panel.pack(fill="both", expand=True)
    ttk.Label(panel, text="Kendi AutoCAD’inizle çizim işleme", font=("Segoe UI", 16, "bold")).pack(anchor="w")
    ttk.Label(panel, text="Çizimlerinizi kaydedip kapatın. İşlem sırasında AutoCAD’i kullanmayın.", wraplength=600).pack(anchor="w", pady=(8, 20))
    origin = tk.StringVar(); code = tk.StringVar(); status = tk.StringVar(value="Bağlantı bekleniyor.")
    try:
        config = load_credentials()
    except Exception:
        config = {}
        status.set("Kayıtlı bağlantı açılamadı. Yeniden eşleştirin.")
    origin.set(config.get("origin", ""))
    ttk.Label(panel, text="Web uygulamasının adresi").pack(anchor="w")
    ttk.Entry(panel, textvariable=origin).pack(fill="x", pady=(4, 10))
    ttk.Label(panel, text="Web uygulamasından aldığınız bağlantı kodu").pack(anchor="w")
    ttk.Entry(panel, textvariable=code, show="•").pack(fill="x", pady=(4, 10))
    worker: list[Worker] = []
    busy = threading.Event()

    def attach(value):
        w = Worker(value, events); worker.append(w)
        threading.Thread(target=w.run, daemon=True).start()

    if config.get("token"):
        attach(config)
        status.set("Bağlantı kayıtlı. Kontrol et ve başlat düğmesini kullanın.")

    def background(fn):
        if busy.is_set():
            return
        busy.set()
        def run():
            try:
                fn()
            except Exception as e:
                events.put(str(e))
            finally:
                busy.clear()
        threading.Thread(target=run, daemon=True).start()

    def pair():
        if worker:
            messagebox.showinfo("ORION", "Yeniden eşleştirmek için bağlantıyı webden kaldırıp yardımcının bağlantı dosyasını sıfırlayın.")
            return
        url, pairing_code = origin.get(), code.get().strip()
        def work():
            value = {"origin": validate_origin(url), "token": secrets.token_hex(32)}
            if not re.fullmatch("[a-f0-9]{64}", pairing_code):
                raise ValueError("Web uygulamasındaki bağlantı kodunun tamamını yapıştırın.")
            api = Api(value)
            result = api.call("pair", {"code": pairing_code, "token": value["token"]}, paired=False)
            value.update(deviceId=result["deviceId"], storageOrigin=result["storageOrigin"])
            save_credentials(value); attach(value)
            events.put("Bağlantı kuruldu. Kontrol et ve başlat düğmesini kullanın.")
        background(work)

    def start():
        if not worker:
            status.set("Önce bilgisayarı web hesabınızla bağlayın."); return
        if worker[0].job:
            status.set("Devam eden çizimin tamamlanmasını bekleyin."); return
        def work():
            state, version, message = autocad_probe(start=True)
            w = worker[0]; w.version = version; w.status(state, message)
            if state == "ready":
                w.active.set()
        background(work)

    buttons = ttk.Frame(panel); buttons.pack(fill="x", pady=8)
    ttk.Button(buttons, text="Bilgisayarı bağla", command=pair).pack(side="left", padx=(0, 8))
    ttk.Button(buttons, text="Kontrol et ve başlat", command=start).pack(side="left", padx=(0, 8))
    def pause():
        if worker:
            worker[0].active.clear()
            if not worker[0].job:
                worker[0].status("attention", "Yeni iş alımı duraklatıldı.")
            else:
                events.put("Bu çizim tamamlandıktan sonra yeni iş alınmayacak.")
    ttk.Button(buttons, text="Yeni iş alımını durdur", command=pause).pack(side="left")
    ttk.Separator(panel).pack(fill="x", pady=16)
    ttk.Label(panel, textvariable=status, wraplength=610, font=("Segoe UI", 11)).pack(anchor="w")
    ttk.Label(panel, text="Üretilen dosyalar bağlantı hatasında bu bilgisayarda korunur.\nPencere açık kalmalıdır; küçültebilirsiniz.", wraplength=610).pack(anchor="w", pady=18)
    ttk.Button(panel, text="Yerel iş dosyalarını aç", command=lambda: os.startfile(str(app_dir()))).pack(anchor="w")
    ttk.Label(panel, text="Sürüm " + VERSION).pack(side="bottom", anchor="e")
    def tick():
        while not events.empty():
            status.set(events.get_nowait())
        root.after(300, tick)
    def close():
        if worker and worker[0].job:
            messagebox.showinfo("ORION", "Çizim sürüyor. Web uygulamasından iptal edebilir veya tamamlanmasını bekleyebilirsiniz."); return
        if worker:
            worker[0].stop.set()
        root.destroy()
    root.protocol("WM_DELETE_WINDOW", close)
    tick(); root.mainloop()
    win32api.CloseHandle(mutex)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
