from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pdfplumber


def scalar(value):
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def inspect_xlsx(path: Path):
    formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
    values = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = []
    for ws, vw in zip(formulas.worksheets, values.worksheets):
        rows = []
        for r_idx, row in enumerate(ws.iter_rows(), 1):
            cells = []
            for c_idx, cell in enumerate(row, 1):
                raw = scalar(cell.value)
                cached = scalar(vw.cell(r_idx, c_idx).value)
                if raw is not None or cached is not None:
                    cells.append({"cell": cell.coordinate, "raw": raw, "value": cached})
            if cells:
                rows.append(cells)
        sheets.append({"title": ws.title, "max_row": ws.max_row, "max_column": ws.max_column, "rows": rows})
    return {"path": str(path), "kind": "xlsx", "sheets": sheets}


def inspect_pdf(path: Path):
    pages = []
    with pdfplumber.open(path) as pdf:
        for idx, page in enumerate(pdf.pages, 1):
            pages.append({"page": idx, "text": page.extract_text(x_tolerance=2, y_tolerance=2) or ""})
    return {"path": str(path), "kind": "pdf", "pages": pages}


def main():
    output = Path(sys.argv[1])
    results = []
    for arg in sys.argv[2:]:
        path = Path(arg)
        results.append(inspect_pdf(path) if path.suffix.lower() == ".pdf" else inspect_xlsx(path))
    output.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print(output)


if __name__ == "__main__":
    main()
